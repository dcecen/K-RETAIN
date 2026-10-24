"""Import de catalogues CSV ou Excel dans la base matériaux.

- CSV : la table est donnée par le nom du fichier (`profiles_*.csv`,
  `palplanches_*.csv`, `butons_*.csv`, `armatures_*.csv`, `aciers_*.csv`,
  `betons_*.csv`) ; séparateur « ; » ou « , » ; décimales « . » ou « , ».
- Excel : une feuille par table, nommée comme la table.
Les en-têtes doivent reprendre les noms de colonnes de la base (unités incluses).

Ligne de commande :
    python -m sout.materiaux.importer <fichier> [<fichier> ...]   # copie dans docs/catalogues puis reconstruit
    python -m sout.materiaux.importer --reconstruire
"""
from __future__ import annotations

import csv
import shutil
import sys
from pathlib import Path

from .base import CATALOGUES, TABLES, reconstruire


def _table_du_nom(nom: str) -> str | None:
    nom = nom.lower()
    for t in TABLES:
        if nom == t or nom.startswith(t + "_") or nom.startswith("demo_" + t):
            return t
    return None


def lire_fichier(chemin: Path) -> list[tuple[str, list[dict]]]:
    """[(table, lignes)] lues dans un fichier CSV ou Excel."""
    if chemin.suffix.lower() == ".csv":
        table = _table_du_nom(chemin.stem)
        if table is None:
            raise ValueError(f"{chemin.name} : le nom doit commencer par une table ({', '.join(TABLES)})")
        texte = chemin.read_text(encoding="utf-8-sig")
        lignes_texte = [l for l in texte.splitlines() if l.strip() and not l.startswith("#")]
        sep = ";" if lignes_texte and ";" in lignes_texte[0] else ","
        return [(table, list(csv.DictReader(lignes_texte, delimiter=sep)))]
    if chemin.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        lots = []
        classeur = load_workbook(chemin, read_only=True, data_only=True)
        for feuille in classeur.worksheets:
            table = _table_du_nom(feuille.title)
            if table is None:
                continue
            lignes = list(feuille.iter_rows(values_only=True))
            if not lignes:
                continue
            entetes = [str(h).strip() if h is not None else "" for h in lignes[0]]
            lots.append((table, [dict(zip(entetes, l)) for l in lignes[1:] if any(v is not None for v in l)]))
        return lots
    raise ValueError(f"{chemin.name} : format non pris en charge (CSV ou XLSX)")


def valider_fichier(chemin: Path) -> dict[str, int]:
    """Lit le fichier et insère toutes ses lignes dans une base en mémoire : toute valeur non numérique
    est signalée (fichier, table, ligne, colonne) avant qu'il ne touche docs/catalogues."""
    import sqlite3

    from .base import creer, inserer

    compte: dict[str, int] = {}
    try:
        c = sqlite3.connect(":memory:")
        creer(c)
        for table, lignes in lire_fichier(chemin):
            compte[table] = compte.get(table, 0) + inserer(c, table, lignes)
        c.close()
    except ValueError as e:
        raise ValueError(f"{chemin.name} : {e}") from None
    if not compte:
        raise ValueError(f"{chemin.name} : aucune ligne reconnue (vérifiez le nom du fichier ou des feuilles et les en-têtes)")
    return compte


def importer(fichiers: list[Path]) -> dict[str, int]:
    """Valide chaque fichier, le copie dans docs/catalogues, puis reconstruit la base."""
    CATALOGUES.mkdir(parents=True, exist_ok=True)
    for f in fichiers:
        valider_fichier(f)
    for f in fichiers:
        if f.resolve().parent != CATALOGUES.resolve():
            shutil.copy2(f, CATALOGUES / f.name)
    return reconstruire()


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(1)
    compte = reconstruire() if args == ["--reconstruire"] else importer([Path(a) for a in args])
    for t, n in compte.items():
        print(f"{t} : {n} ligne(s)")
