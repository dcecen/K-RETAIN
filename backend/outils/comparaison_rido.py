"""Dossier de comparaison K-Retain / RIDO : un classeur Excel prêt à remplir avec les résultats RIDO.

    cd backend
    .venv/Scripts/python.exe -m outils.comparaison_rido                 # dossier exemples/rido
    .venv/Scripts/python.exe -m outils.comparaison_rido chemin/dossier

Pour chaque fichier RIDO du dossier (*.rid, *.rido, *.dat), le script K-Retain converti est écrit à côté
(<nom>.yaml) avec son rapport de conversion (<nom>.rapport.txt). Chaque script <nom>.yaml du dossier est
ensuite calculé et le classeur <dossier>/comparaison_rido.xlsx est écrit :

- feuille « Synthèse » : une ligne par cas, phase et grandeur (déplacement, moments, tranchant, efforts
  d'appui), la valeur K-Retain, une colonne RIDO à saisir, l'écart et le verdict (tolérance 5 %, CLAUDE.md
  règle 3) calculés par des formules Excel ;
- une feuille par cas : profils z, y, M, V de chaque phase, pour tracer les courbes face à celles de RIDO.

Si un fichier <nom>.resultats_rido.yaml existe (liste de {phase, grandeur, valeur}), la colonne RIDO est
pré-remplie avec ses valeurs. Aucune valeur RIDO n'est inventée : sans ce fichier, la colonne reste vide.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from ruamel.yaml import YAML

RACINE = Path(__file__).resolve().parents[2]
DOSSIER = RACINE / "exemples" / "rido"
EXTENSIONS_RIDO = (".rid", ".rido", ".dat")
TOLERANCE = 0.05   # CLAUDE.md, règle 3 : tolérance des cas de validation

CONVENTIONS = ("Conventions K-Retain : y positif vers la fouille (mm) ; M positif quand la face côté fouille est tendue "
               "(kN·m/m) ; V = dM/dz, z vers le haut (kN/m) ; effort d'appui positif quand l'appui retient l'écran "
               "(kN/m). Vérifier les conventions de signe de RIDO avant de comparer.")


def convertir_fichiers_rido(dossier: Path) -> list[Path]:
    """Convertit chaque fichier RIDO du dossier en script K-Retain ; renvoie les scripts écrits."""
    from sout.script.rido import convertir

    ecrits = []
    for f in sorted(p for p in dossier.iterdir() if p.suffix.lower() in EXTENSIONS_RIDO):
        conv = convertir(f.read_text(encoding="utf-8", errors="replace"), f.name)
        f.with_suffix(".rapport.txt").write_text("\n".join(conv.rapport) + "\n", encoding="utf-8")
        if conv.repris:
            cible = f.with_suffix(".yaml")
            cible.write_text(conv.texte, encoding="utf-8")
            ecrits.append(cible)
    return ecrits


def calculer(texte: str) -> dict:
    from fastapi.testclient import TestClient

    from sout.api import app

    r = TestClient(app).post("/api/calcul", json={"texte": texte}).json()
    if not r.get("ok"):
        raise RuntimeError("; ".join(str(e.get("message", e)) for e in r.get("erreurs", [])) or "calcul refusé")
    return r["resultats"]


def grandeurs(resultats: dict) -> list[tuple[int, str, str, str, float | None]]:
    """(numéro de phase, phase, grandeur, unité, valeur K-Retain) pour chaque phase calculée."""
    lignes = []
    for i, ph in enumerate(resultats["phases"]):
        nom = ph["nom"]
        if not ph.get("ok"):
            lignes.append((i, nom, "Phase en rupture : " + str(ph.get("message") or ""), "", None))
            continue
        s = ph["synthese"]
        lignes += [
            (i, nom, "Déplacement en tête", "mm", s["y_tete_mm"]),
            (i, nom, "Déplacement maximal", "mm", s["y_max_mm"]),
            (i, nom, "Cote du déplacement maximal", "m", s["z_y_max"]),
            (i, nom, "Moment maximal (M > 0)", "kN·m/m", s["M_max"]),
            (i, nom, "Moment minimal (M < 0)", "kN·m/m", s["M_min"]),
            (i, nom, "Tranchant maximal |V|", "kN/m", s["V_absmax"]),
        ]
        for a in ph.get("appuis") or []:
            if a.get("actif", True):
                lignes.append((i, nom, f"Effort d'appui {a['nom']} (cote {a['cote']:g} m)", "kN/m", a["effort_kN_m"]))
    return lignes


def _lire_resultats_rido(chemin: Path) -> dict[tuple[str, str], float]:
    if not chemin.exists():
        return {}
    donnees = YAML(typ="safe").load(chemin.read_text(encoding="utf-8")) or []
    return {(str(d["phase"]), str(d["grandeur"])): float(d["valeur"]) for d in donnees}


def ecrire_classeur(cas: list[tuple[str, dict | str]], sortie: Path, rido: dict[str, dict] | None = None) -> Path:
    """cas : (nom, résultats ou message d'erreur). rido : valeurs RIDO connues par cas, clé (phase, grandeur)."""
    rido = rido or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    gras, titre = Font(bold=True), Font(bold=True, size=13)
    a_saisir = PatternFill("solid", fgColor="FFF4CC")
    ws["A1"] = "Comparaison K-Retain / RIDO"
    ws["A1"].font = titre
    ws["A2"] = CONVENTIONS
    ws["A3"] = (f"Colonne RIDO (fond jaune) à saisir. Écart = (K-Retain − RIDO) / |RIDO| ; verdict « conforme » "
                f"si |écart| ≤ {TOLERANCE:.0%}.")
    entetes = ["Cas", "N° phase", "Phase", "Grandeur", "Unité", "K-Retain", "RIDO", "Écart", "Verdict"]
    for j, e in enumerate(entetes, 1):
        ws.cell(5, j, e).font = gras
    ligne = 6
    for nom, res in cas:
        if isinstance(res, str):
            ws.cell(ligne, 1, nom)
            ws.cell(ligne, 4, "Calcul impossible : " + res)
            ligne += 1
            continue
        connues = rido.get(nom, {})
        for i, phase, g, unite, v in grandeurs(res):
            ws.cell(ligne, 1, nom)
            ws.cell(ligne, 2, i)
            ws.cell(ligne, 3, phase)
            ws.cell(ligne, 4, g)
            ws.cell(ligne, 5, unite)
            if v is not None:
                ws.cell(ligne, 6, round(v, 4 if unite == "mm" else 2))
                c = ws.cell(ligne, 7, connues.get((phase, g)))
                c.fill = a_saisir
                ws.cell(ligne, 8, f'=IF(OR(G{ligne}="",G{ligne}=0),"",(F{ligne}-G{ligne})/ABS(G{ligne}))').number_format = "0.0%"
                ws.cell(ligne, 9, f'=IF(H{ligne}="","",IF(ABS(H{ligne})<={TOLERANCE},"conforme","ÉCART"))')
            ligne += 1
    for j, l in enumerate([26, 9, 34, 38, 9, 12, 12, 10, 11], 1):
        ws.column_dimensions[get_column_letter(j)].width = l
    ws.freeze_panes = "A6"
    for nom, res in cas:
        if not isinstance(res, str):
            _feuille_profils(wb, nom, res)
    sortie.parent.mkdir(parents=True, exist_ok=True)
    wb.save(sortie)
    return sortie


def _feuille_profils(wb: Workbook, nom: str, res: dict) -> None:
    ws = wb.create_sheet(nom[:31])
    ws["A1"] = f"Profils par phase : {res.get('projet', nom)}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = CONVENTIONS
    col = 1
    for i, ph in enumerate(res["phases"]):
        if not ph.get("ok"):
            continue
        ws.cell(4, col, f"Phase {i} : {ph['nom']}").font = Font(bold=True)
        ws.cell(4, col).alignment = Alignment(wrap_text=False)
        for j, e in enumerate(["z (m)", "y (mm)", "z diag (m)", "M (kN·m/m)", "V (kN/m)"]):
            ws.cell(5, col + j, e).font = Font(bold=True)
        for k, (z, y) in enumerate(zip(ph["z"], ph["y_mm"], strict=False)):
            ws.cell(6 + k, col, round(z, 3))
            ws.cell(6 + k, col + 1, round(y, 4))
        for k, (z, m, v) in enumerate(zip(ph["z_diag"], ph["M"], ph["V"], strict=False)):
            ws.cell(6 + k, col + 2, round(z, 3))
            ws.cell(6 + k, col + 3, round(m, 2))
            ws.cell(6 + k, col + 4, round(v, 2))
        col += 6
    ws.freeze_panes = "A6"


def comparer(dossier: Path = DOSSIER) -> Path:
    convertir_fichiers_rido(dossier)
    cas: list[tuple[str, dict | str]] = []
    rido: dict[str, dict] = {}
    for script in sorted(dossier.glob("*.yaml")):
        if script.name.endswith(".resultats_rido.yaml"):
            continue
        nom = script.stem
        try:
            cas.append((nom, calculer(script.read_text(encoding="utf-8"))))
        except Exception as e:  # noqa: BLE001 - le cas en échec est signalé dans le classeur, les autres continuent
            cas.append((nom, str(e)))
        rido[nom] = _lire_resultats_rido(dossier / f"{nom}.resultats_rido.yaml")
    return ecrire_classeur(cas, dossier / "comparaison_rido.xlsx", rido)


if __name__ == "__main__":
    print(comparer(Path(sys.argv[1]) if len(sys.argv) > 1 else DOSSIER))
