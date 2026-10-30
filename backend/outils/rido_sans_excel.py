"""Version d'un fichier RIDO indépendante de sa feuille Excel (« Données RIDO.xls »).

    cd backend
    .venv/Scripts/python.exe -m outils.rido_sans_excel fichier.rid "Données RIDO.xlsx"
    .venv/Scripts/python.exe -m outils.rido_sans_excel fichier.rid           # gabarit à remplir, sans classeur

Avec le classeur : chaque xlc(feuille, ligne, colonne) est remplacé par la valeur de la cellule (les libellés
l_X et c_X de la directive CONST sont lus dans la feuille 1 : libellés de ligne en colonnes 1 et 2, libellés de
colonne en ligne 1), les directives « # XLS= » et « # CONST= » sont retirées ; tout le reste du fichier (formules
#, commandes) est gardé tel quel. Écrit <fichier>_sans_excel.rid à côté du fichier, en Windows-1252 (RIDO).

Sans classeur : chaque cellule lue devient une variable #X<ligne>_<colonne> = A_RENSEIGNER déclarée en tête,
avec la cellule d'origine en commentaire : il reste à recopier les valeurs.

Classeur .xlsx lu par openpyxl ; .xls (ancien format) par xlrd s'il est installé, sinon enregistrez la feuille
en .xlsx depuis Excel.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from sout.script.rido import Classeur, ErreurExpression, ValeurIndisponible, evaluer

XLC = re.compile(r"xlc\(([^()]*)\)", re.IGNORECASE)
DIRECTIVE = re.compile(r"^\s*#\s*(XLS|CONST)\s*=", re.IGNORECASE)


def lire_texte(chemin: Path) -> str:
    b = chemin.read_bytes()
    try:
        return b.decode("utf-8")
    except UnicodeDecodeError:
        return b.decode("cp1252", errors="replace")


def ecrire_rido(chemin: Path, texte: str) -> None:
    """Fichier RIDO : Windows-1252 et fins de ligne CRLF (logiciel Windows)."""
    chemin.write_bytes(texte.replace("\r\n", "\n").replace("\n", "\r\n").encode("cp1252", errors="replace"))


def charger_classeur(chemin: Path) -> Classeur:
    cellules: dict[tuple[int, int, int], object] = {}
    if chemin.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError:
            raise SystemExit("Lecture d'un .xls : installez xlrd (pip install xlrd) ou enregistrez la feuille en .xlsx") from None
        wb = xlrd.open_workbook(str(chemin))
        for f in range(wb.nsheets):
            ws = wb.sheet_by_index(f)
            for l in range(ws.nrows):
                for c in range(ws.ncols):
                    v = ws.cell_value(l, c)
                    if v not in ("", None):
                        cellules[(f + 1, l + 1, c + 1)] = v
    else:
        from openpyxl import load_workbook
        wb = load_workbook(chemin, data_only=True, read_only=True)
        for f, ws in enumerate(wb.worksheets, 1):
            for ligne in ws.iter_rows():
                for cel in ligne:
                    if cel.value not in ("", None) and hasattr(cel, "row"):
                        cellules[(f, cel.row, cel.column)] = cel.value
    return Classeur(cellules, chemin.name)


def _format(v: float) -> str:
    return f"{v:.10g}"


def sans_excel(texte: str, classeur: Classeur | None, nom_classeur: str = "") -> tuple[str, list[str]]:
    """(texte RIDO autonome, messages). Avec classeur : valeurs recopiées ; sans : gabarit A_RENSEIGNER."""
    messages: list[str] = []
    cellules: dict[str, str] = {}

    def remplacer(m: re.Match) -> str:
        brut = m.group(0)
        try:
            args = [evaluer(a, {}, classeur) for a in m.group(1).split(",")]
            if len(args) != 3:
                raise ErreurExpression("xlc attend 3 arguments")
        except (ErreurExpression, ValeurIndisponible) as e:
            if classeur is None:
                nom = "X_" + re.sub(r"\W", "_", m.group(1))
                cellules[nom] = brut
                return nom
            messages.append(f"{brut} : {e}")
            return brut
        if classeur is None:
            nom = f"X{int(args[1])}_{int(args[2])}" + (f"_F{int(args[0])}" if int(args[0]) != 1 else "")
            cellules[nom] = brut
            return nom
        try:
            return _format(classeur.xlc(*args))
        except ErreurExpression as e:
            messages.append(str(e))
            return brut

    lignes = texte.splitlines()
    sortie: list[str] = []
    for i, l in enumerate(lignes):
        if i > 0 and DIRECTIVE.match(l):
            continue
        sortie.append(l if (i == 0 or l.lstrip().startswith("*")) else XLC.sub(remplacer, l))
    if classeur is not None:
        entete = [f"*Version indépendante de la feuille Excel {nom_classeur} : valeurs des cellules recopiées"]
    else:
        entete = ["*Version indépendante de la feuille Excel : recopiez chaque valeur A_RENSEIGNER"]
        for nom, brut in sorted(cellules.items(), key=lambda kv: [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", kv[0])]):
            entete += [f"*{nom} : cellule {brut}", f"#{nom} = A_RENSEIGNER"]
    return "\n".join(sortie[:1] + entete + sortie[1:]) + "\n", messages


def main(argv: list[str]) -> None:
    if not argv:
        raise SystemExit(__doc__)
    source = Path(argv[0])
    classeur = charger_classeur(Path(argv[1])) if len(argv) > 1 else None
    texte, messages = sans_excel(lire_texte(source), classeur, Path(argv[1]).name if len(argv) > 1 else "")
    cible = source.with_name(source.stem.removesuffix("_excel") + "_sans_excel.rid")
    ecrire_rido(cible, texte)
    print(cible)
    for m in messages:
        print("A verifier :", m)


if __name__ == "__main__":
    main(sys.argv[1:])
