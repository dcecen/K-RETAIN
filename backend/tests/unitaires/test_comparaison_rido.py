"""Outil de comparaison K-Retain / RIDO (outils/comparaison_rido.py) : classeur, valeurs reprises, formules."""
from pathlib import Path

from openpyxl import load_workbook

from outils.comparaison_rido import calculer, ecrire_classeur, grandeurs

RACINE = Path(__file__).resolve().parents[3]


def test_classeur_reprend_les_resultats_k_retain(tmp_path):
    res = calculer((RACINE / "exemples" / "fouille_exemple.yaml").read_text(encoding="utf-8"))
    lignes = grandeurs(res)
    finale = len(res["phases"]) - 1
    ph = res["phases"][finale]
    y_max = next(v for i, _, g, _, v in lignes if i == finale and g == "Déplacement maximal")
    assert y_max == ph["synthese"]["y_max_mm"]
    appuis = [g for i, _, g, _, _ in lignes if i == finale and g.startswith("Effort d'appui")]
    assert len(appuis) == sum(1 for a in ph["appuis"] if a.get("actif", True))

    # Une valeur RIDO connue est pré-remplie, les autres restent vides ; l'écart est une formule Excel.
    cle = (ph["nom"], "Déplacement maximal")
    sortie = ecrire_classeur([("fouille_exemple", res), ("cas_en_erreur", "script invalide")],
                             tmp_path / "c.xlsx", {"fouille_exemple": {cle: 10.0}})
    ws = load_workbook(sortie)["Synthèse"]
    rangs = [r for r in ws.iter_rows(min_row=6, values_only=True) if r[0] == "fouille_exemple"]
    remplies = [r for r in rangs if r[6] is not None]
    assert len(remplies) == 1 and remplies[0][2] == ph["nom"] and remplies[0][6] == 10.0
    assert str(remplies[0][7]).startswith("=IF(")
    assert any(r[0] == "cas_en_erreur" and "script invalide" in r[3] for r in ws.iter_rows(min_row=6, values_only=True))
    assert "fouille_exemple" in load_workbook(sortie).sheetnames
